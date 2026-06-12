# -*- coding: utf-8 -*-
"""Генератор презентации «Битрикс24 КЭДО» в фирменном стиле Битрикс24
для партнёра Smart Business."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Фирменная палитра Битрикс24 ----
BLUE = RGBColor(0x00, 0x75, 0xFF)      # основной синий
CYAN = RGBColor(0x2F, 0xC6, 0xF5)      # фирменный голубой
NAVY = RGBColor(0x0B, 0x1F, 0x3A)      # тёмно-синий фон
DARK = RGBColor(0x1E, 0x2A, 0x38)      # заголовки
GRAY = RGBColor(0x53, 0x5C, 0x69)      # основной текст (фирменный серый Б24)
LIGHT = RGBColor(0xF0, 0xF6, 0xFC)     # светлый фон карточек
GREEN = RGBColor(0x9D, 0xCF, 0x00)     # фирменный зелёный акцент
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xF9, 0xA9, 0x1D)

FONT = "Arial"
FONT_B = "Arial"

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def rect(slide, x, y, w, h, color, rounded=False, radius=0.12):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    _solid(shp, color)
    return shp


def text(slide, x, y, w, h, runs, size=18, color=GRAY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
         space_after=6, font=FONT):
    """runs: str | list[str] | list[(str, dict)]"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [runs]
    for i, item in enumerate(runs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line_spacing
        para.space_after = Pt(space_after)
        if isinstance(item, tuple):
            txt, opts = item
        else:
            txt, opts = item, {}
        r = para.add_run()
        r.text = txt
        f = r.font
        f.name = opts.get("font", font)
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.color.rgb = opts.get("color", color)
    return tb


def bullet_block(slide, x, y, w, items, size=15, gap=Inches(0.62),
                 color=GRAY, dot_color=BLUE):
    """Маркированный список с фирменными точками."""
    cy = y
    for head, body in items:
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, cy + Inches(0.09),
                                   Inches(0.13), Inches(0.13))
        _solid(d, dot_color)
        runs = [(head, {"bold": True, "color": DARK, "size": size})]
        if body:
            runs.append((body, {"size": size - 1, "color": color}))
        text(slide, x + Inches(0.3), cy - Inches(0.05), w - Inches(0.3),
             gap, runs, space_after=2, line_spacing=1.0)
        cy += gap


def sb_logo(slide, x, y, scale=1.0, dark_bg=False):
    """Логотип Smart Business: фирменный знак + текст."""
    s = scale
    box = rect(slide, x, y, Inches(0.42 * s), Inches(0.42 * s), BLUE,
               rounded=True, radius=0.28)
    text(slide, x, y - Inches(0.02 * s), Inches(0.42 * s), Inches(0.42 * s),
         [("S", {"size": int(20 * s), "bold": True, "color": WHITE})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + Inches(0.5 * s), y - Inches(0.06 * s), Inches(2.6 * s),
         Inches(0.5 * s),
         [("smart ", {"size": int(15 * s), "bold": True,
                      "color": WHITE if dark_bg else DARK})],
         anchor=MSO_ANCHOR.MIDDLE)
    # второе слово отдельным цветом — добавим в тот же textbox нельзя, проще два бокса
    text(slide, x + Inches(1.08 * s), y - Inches(0.06 * s), Inches(2.2 * s),
         Inches(0.5 * s),
         [("business", {"size": int(15 * s), "bold": True, "color": CYAN})],
         anchor=MSO_ANCHOR.MIDDLE)


def header(slide, title, subtitle=None, num=None, total=12):
    """Шапка контентного слайда."""
    rect(slide, 0, 0, SW, Inches(0.12), BLUE)          # фирменная полоса
    rect(slide, 0, 0, Inches(4.2), Inches(0.12), CYAN)
    text(slide, Inches(0.55), Inches(0.42), Inches(11), Inches(0.8),
         [(title, {"size": 30, "bold": True, "color": DARK})])
    if subtitle:
        text(slide, Inches(0.55), Inches(1.08), Inches(11.5), Inches(0.5),
             [(subtitle, {"size": 15, "color": GRAY})])
    sb_logo(slide, Inches(10.55), Inches(0.45), scale=0.85)
    if num:
        text(slide, Inches(12.45), Inches(7.05), Inches(0.7), Inches(0.35),
             [(f"{num:02d}", {"size": 11, "color": GRAY, "bold": True})],
             align=PP_ALIGN.RIGHT)
    text(slide, Inches(0.55), Inches(7.05), Inches(6), Inches(0.35),
         [("Битрикс24 КЭДО  •  Smart Business — официальный партнёр Битрикс24",
           {"size": 10, "color": GRAY})])


def card(slide, x, y, w, h, head, body, head_color=DARK, accent=BLUE,
         head_size=16, body_size=13, fill=LIGHT):
    c = rect(slide, x, y, w, h, fill, rounded=True, radius=0.08)
    rect(slide, x, y, Inches(0.09), h, accent, rounded=False)
    text(slide, x + Inches(0.28), y + Inches(0.16), w - Inches(0.45),
         Inches(0.6), [(head, {"size": head_size, "bold": True,
                               "color": head_color})], space_after=2)
    text(slide, x + Inches(0.28), y + Inches(0.62), w - Inches(0.45),
         h - Inches(0.75), [(body, {"size": body_size, "color": GRAY})],
         line_spacing=1.05)
    return c


def num_card(slide, x, y, w, h, big, label, color=BLUE):
    rect(slide, x, y, w, h, LIGHT, rounded=True, radius=0.1)
    text(slide, x, y + Inches(0.18), w, Inches(0.85),
         [(big, {"size": 40, "bold": True, "color": color})],
         align=PP_ALIGN.CENTER)
    text(slide, x + Inches(0.2), y + Inches(1.05), w - Inches(0.4),
         h - Inches(1.1), [(label, {"size": 13, "color": GRAY})],
         align=PP_ALIGN.CENTER, line_spacing=1.0)


def step_chip(slide, x, y, w, h, n, head, body, color=BLUE):
    rect(slide, x, y, w, h, LIGHT, rounded=True, radius=0.12)
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.22),
                               y + Inches(0.2), Inches(0.5), Inches(0.5))
    _solid(c, color)
    text(slide, x + Inches(0.22), y + Inches(0.18), Inches(0.5), Inches(0.5),
         [(str(n), {"size": 18, "bold": True, "color": WHITE})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + Inches(0.88), y + Inches(0.14), w - Inches(1.05),
         Inches(0.62), [(head, {"size": 15, "bold": True, "color": DARK})],
         line_spacing=0.95)
    text(slide, x + Inches(0.22), y + Inches(0.8), w - Inches(0.44),
         h - Inches(0.92), [(body, {"size": 12, "color": GRAY})],
         line_spacing=1.0)


# ============================================================ 1. Титул
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(7.32), SW, Inches(0.18), CYAN)
# декоративные круги в стиле Б24
for cx, cy, d, col in [(11.2, -1.2, 4.2, BLUE), (12.3, 5.6, 3.4, CYAN),
                       (-1.6, 5.2, 3.6, RGBColor(0x14, 0x2E, 0x52))]:
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy),
                           Inches(d), Inches(d))
    _solid(o, col)
    o.fill.fore_color.rgb = col
    o.fill.transparency = 0  # pptx не поддерживает прозрачность напрямую
sb_logo(s, Inches(0.7), Inches(0.6), scale=1.2, dark_bg=True)
text(s, Inches(0.7), Inches(2.35), Inches(10.5), Inches(1.0),
     [("Битрикс24 ", {"size": 54, "bold": True, "color": WHITE})])
text(s, Inches(4.05), Inches(2.35), Inches(6), Inches(1.0),
     [("КЭДО", {"size": 54, "bold": True, "color": CYAN})])
text(s, Inches(0.7), Inches(3.55), Inches(9.6), Inches(1.2),
     [("Кадровый электронный документооборот с Госключом:",
       {"size": 22, "color": WHITE}),
      ("документы подписываются за минуты — без бумаги, печати и курьеров",
       {"size": 18, "color": RGBColor(0xB9, 0xC9, 0xDE)})],
     line_spacing=1.1)
rect(s, Inches(0.7), Inches(5.35), Inches(4.6), Inches(0.66), BLUE,
     rounded=True, radius=0.5)
text(s, Inches(0.7), Inches(5.35), Inches(4.6), Inches(0.66),
     [("Официальный партнёр Битрикс24", {"size": 16, "bold": True,
                                         "color": WHITE})],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(0.7), Inches(6.35), Inches(9), Inches(0.5),
     [("Smart Business  •  внедрение, настройка и сопровождение Битрикс24",
       {"size": 14, "color": RGBColor(0xB9, 0xC9, 0xDE)})])

# ============================================================ 2. Что такое КЭДО
s = prs.slides.add_slide(BLANK)
header(s, "Что такое КЭДО", "Подписание кадровых документов между компанией и "
       "сотрудником онлайн — без бумаги и личных встреч", 2)
card(s, Inches(0.55), Inches(1.8), Inches(6.0), Inches(2.3),
     "Полностью юридически значимо",
     "Электронные кадровые документы равнозначны бумажным. Правовая база:\n"
     "• Федеральный закон № 377-ФЗ от 22.11.2021\n"
     "• Статьи 22.1–22.3 Трудового кодекса РФ\n"
     "Дублировать документы на бумаге не нужно.", accent=BLUE)
card(s, Inches(0.55), Inches(4.35), Inches(6.0), Inches(2.3),
     "Где находится в Битрикс24",
     "Раздел «КЭДО + Госключ» в главном меню портала.\n"
     "Сотрудники подписывают документы из браузера или мобильного "
     "приложения Битрикс24 — где бы они ни находились: в офисе, "
     "на удалёнке или в командировке.", accent=CYAN)
card(s, Inches(6.85), Inches(1.8), Inches(5.95), Inches(2.3),
     "Для кого",
     "• Компании с удалёнными сотрудниками и филиалами\n"
     "• HR-отделы с большим потоком документов\n"
     "• Бизнес, который уже работает в Битрикс24 и 1С\n"
     "• Компании от 10 до 10 000+ сотрудников", accent=GREEN)
card(s, Inches(6.85), Inches(4.35), Inches(5.95), Inches(2.3),
     "Что переводится в электронный вид",
     "Трудовые договоры и допсоглашения, приказы, заявления "
     "(отпуск, командировка, отгул), согласия на обработку персональных "
     "данных, положения и регламенты, ознакомление с ЛНА.", accent=ORANGE)

# ============================================================ 3. Почему пора
s = prs.slides.add_slide(BLANK)
header(s, "Почему бумажный кадровый документооборот — это дорого", None, 3)
bullet_block(s, Inches(0.55), Inches(1.7), Inches(5.9), [
    ("Долго", " — подписание с удалённым сотрудником занимает дни: печать, "
     "курьер, ожидание, возврат."),
    ("Дорого", " — бумага, печать, доставка, архивное хранение, время HR."),
    ("Рискованно", " — документы теряются, подписываются с опозданием или "
     "не теми датами."),
    ("Непрозрачно", " — непонятно, на каком этапе документ и кто его "
     "задерживает."),
], gap=Inches(0.95), size=16)
num_card(s, Inches(6.85), Inches(1.7), Inches(2.85), Inches(2.2),
         "−80%", "времени HR-отдела на создание, согласование и "
         "подписание документов", BLUE)
num_card(s, Inches(9.9), Inches(1.7), Inches(2.85), Inches(2.2),
         "30 ч", "в месяц экономит кадровый специалист с Битрикс24 КЭДО",
         CYAN)
num_card(s, Inches(6.85), Inches(4.1), Inches(2.85), Inches(2.2),
         "2 мин", "среднее время подписания документа сотрудником через "
         "Госключ", GREEN)
num_card(s, Inches(9.9), Inches(4.1), Inches(2.85), Inches(2.2),
         "0 ₽", "расходов на бумагу, печать, курьеров и физический архив",
         ORANGE)

# ============================================================ 4. Возможности
s = prs.slides.add_slide(BLANK)
header(s, "Возможности Битрикс24 КЭДО",
       "Полный цикл работы с кадровыми документами в одном окне", 4)
feats = [
    ("Шаблоны документов", "Готовые и собственные шаблоны: данные компании "
     "и сотрудника подставляются автоматически из Битрикс24.", BLUE),
    ("Комплекты документов", "Наборы для типовых ситуаций: приём сотрудника "
     "— трудовой договор, NDA, согласие на обработку ПД одним кликом.", CYAN),
    ("Маршруты подписания", "Согласующие по ролям, а не по фамилиям: "
     "«руководитель отдела», «директор». Маршрут — от компании или от "
     "сотрудника.", GREEN),
    ("Подписание через Госключ", "Бесплатное мобильное приложение: сотрудник "
     "подписывает документ со смартфона за пару минут.", BLUE),
    ("Интеграция с 1С", "Отправка документов из 1С на подпись и "
     "автозаполнение данных сотрудников — без ручного ввода.", CYAN),
    ("Электронный архив", "Все документы хранятся в одном месте: даже через "
     "5 лет нужный файл находится в 2 клика.", GREEN),
]
for i, (h, b, a) in enumerate(feats):
    x = Inches(0.55 + (i % 3) * 4.18)
    y = Inches(1.75 + (i // 3) * 2.5)
    card(s, x, y, Inches(3.95), Inches(2.3), h, b, accent=a)

# ============================================================ 5. Как работает
s = prs.slides.add_slide(BLANK)
header(s, "Как отправить документ на подпись",
       "Четыре шага на стороне HR — и документ уходит сотрудникам", 5)
steps = [
    ("Подготовьте документ", "Загрузите файл или создайте его по шаблону — "
     "реквизиты заполнятся автоматически."),
    ("Выберите компанию и представителя", "Кто подписывает со стороны "
     "работодателя (УКЭП)."),
    ("Укажите провайдера подписи", "Госключ (УНЭП/УКЭП) или внутренний "
     "сервис Битрикс24 КЭДО (ПЭП)."),
    ("Выберите сотрудников", "Одного, отдел или всю компанию — каждый "
     "получит уведомление."),
]
for i, (h, b) in enumerate(steps):
    step_chip(s, Inches(0.55 + i * 3.13), Inches(1.85), Inches(2.95),
              Inches(2.0), i + 1, h, b)
rect(s, Inches(0.55), Inches(4.3), Inches(12.23), Inches(2.25), NAVY,
     rounded=True, radius=0.07)
text(s, Inches(0.95), Inches(4.55), Inches(11.5), Inches(0.5),
     [("Что видит сотрудник", {"size": 17, "bold": True, "color": CYAN})])
text(s, Inches(0.95), Inches(5.1), Inches(11.5), Inches(1.3),
     [("1.  Уведомление в Битрикс24 о поступившем документе",
       {"size": 14, "color": WHITE}),
      ("2.  Переход в приложение Госключ на смартфоне",
       {"size": 14, "color": WHITE}),
      ("3.  Подписание в пару касаний — документ автоматически возвращается "
       "в Битрикс24 и попадает в архив", {"size": 14, "color": WHITE})],
     line_spacing=1.15)

# ============================================================ 6. Подписи и Госключ
s = prs.slides.add_slide(BLANK)
header(s, "Электронные подписи и Госключ",
       "Каждая сторона подписывает документы так, как требует закон", 6)
card(s, Inches(0.55), Inches(1.8), Inches(6.0), Inches(2.15),
     "Работодатель — УКЭП",
     "Усиленная квалифицированная электронная подпись. Документы заверяет "
     "уполномоченный представитель компании.", accent=BLUE)
card(s, Inches(0.55), Inches(4.2), Inches(6.0), Inches(2.15),
     "Сотрудник — УНЭП / ПЭП",
     "Усиленная неквалифицированная подпись через Госключ или простая "
     "электронная подпись через внутренний сервис Битрикс24 КЭДО — "
     "в зависимости от типа документа.", accent=CYAN)
rect(s, Inches(6.85), Inches(1.8), Inches(5.95), Inches(4.55), NAVY,
     rounded=True, radius=0.06)
text(s, Inches(7.25), Inches(2.1), Inches(5.2), Inches(0.5),
     [("Госключ", {"size": 20, "bold": True, "color": CYAN})])
text(s, Inches(7.25), Inches(2.65), Inches(5.2), Inches(3.5),
     [("•  Бесплатное мобильное приложение Минцифры",
       {"size": 14, "color": WHITE}),
      ("•  Подпись выпускается онлайн — без визита в удостоверяющий центр",
       {"size": 14, "color": WHITE}),
      ("•  Нужна только подтверждённая учётная запись на Госуслугах",
       {"size": 14, "color": WHITE}),
      ("•  Сотрудник подписывает документы со своего смартфона из любой "
       "точки", {"size": 14, "color": WHITE}),
      ("•  Для сотрудника — 0 рублей", {"size": 14, "bold": True,
                                        "color": GREEN})],
     line_spacing=1.25)

# ============================================================ 7. Интеграция с 1С
s = prs.slides.add_slide(BLANK)
header(s, "Интеграция с 1С: без ручного ввода",
       "Связка Битрикс24 КЭДО + 1С (в т.ч. 1С:ЗУП) исключает ошибки и "
       "двойную работу", 7)
card(s, Inches(0.55), Inches(1.85), Inches(3.95), Inches(2.6),
     "Документы из 1С — на подпись",
     "Кадровик формирует документ в привычной 1С и отправляет его на "
     "подписание в Битрикс24 КЭДО, не выходя из программы.", accent=BLUE)
card(s, Inches(4.7), Inches(1.85), Inches(3.95), Inches(2.6),
     "Данные сотрудников из 1С",
     "Битрикс24 КЭДО заполняет шаблоны данными из 1С: сопоставление "
     "сотрудников настраивается один раз.", accent=CYAN)
card(s, Inches(8.85), Inches(1.85), Inches(3.95), Inches(2.6),
     "Статусы и архив",
     "Подписанные документы и статусы синхронизируются: в 1С видно, что "
     "подписано, а что ждёт сотрудника.", accent=GREEN)
rect(s, Inches(0.55), Inches(4.8), Inches(12.25), Inches(1.7), LIGHT,
     rounded=True, radius=0.09)
text(s, Inches(0.95), Inches(5.0), Inches(11.5), Inches(1.3),
     [("Результат: ", {"size": 16, "bold": True, "color": DARK}),
      ("ноль ручного переноса данных между системами, ноль ошибок в "
       "реквизитах и датах, единый контур «1С → Битрикс24 → Госключ → "
       "архив».", {"size": 16, "color": GRAY})],
     line_spacing=1.15)

# ============================================================ 8. Шаблоны и комплекты
s = prs.slides.add_slide(BLANK)
header(s, "Шаблоны и комплекты документов",
       "Типовые документы создаются за секунды, а не печатаются заново", 8)
bullet_block(s, Inches(0.55), Inches(1.8), Inches(6.0), [
    ("Свои шаблоны", " — загрузите трудовой договор, допсоглашение, приказ "
     "или заявление компании."),
    ("Автозаполнение", " — данные компании и сотрудника подставляются из "
     "Битрикс24 и 1С."),
    ("Самообслуживание", " — сотрудник сам формирует заявление на отпуск "
     "или командировку по шаблону и отправляет на подпись."),
    ("Подпись ПЭП", " — типовые заявления подписываются простой электронной "
     "подписью прямо в Битрикс24."),
], gap=Inches(1.05), size=15)
rect(s, Inches(6.85), Inches(1.8), Inches(5.95), Inches(4.6), LIGHT,
     rounded=True, radius=0.06)
text(s, Inches(7.2), Inches(2.05), Inches(5.3), Inches(0.5),
     [("Комплект «Приём сотрудника»", {"size": 17, "bold": True,
                                       "color": DARK})])
for i, doc in enumerate(["Трудовой договор", "Согласие на обработку "
                         "персональных данных", "Положение о коммерческой "
                         "тайне (NDA)", "Ознакомление с ЛНА и должностной "
                         "инструкцией"]):
    y = Inches(2.65 + i * 0.78)
    chk = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.2), y, Inches(0.32),
                             Inches(0.32))
    _solid(chk, GREEN)
    text(s, Inches(7.2), y - Inches(0.04), Inches(0.32), Inches(0.36),
         [("✓", {"size": 14, "bold": True, "color": WHITE})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(7.68), y - Inches(0.05), Inches(5.0), Inches(0.7),
         [(doc, {"size": 14, "color": GRAY})], line_spacing=0.95)
text(s, Inches(7.2), Inches(5.85), Inches(5.3), Inches(0.5),
     [("Весь комплект уходит сотруднику одним кликом",
       {"size": 14, "bold": True, "color": BLUE})])

# ============================================================ 9. Ограничения закона
s = prs.slides.add_slide(BLANK)
header(s, "Что остаётся на бумаге", "Закон (ст. 22.1 ТК РФ) оставляет "
       "несколько исключений — важно учесть при переходе", 9)
exc = [
    "Трудовые книжки и сведения о трудовой деятельности (СТД-Р)",
    "Акты о несчастном случае на производстве",
    "Приказы (распоряжения) об увольнении",
    "Документы о прохождении инструктажей по охране труда",
]
for i, e in enumerate(exc):
    y = Inches(1.85 + i * 0.95)
    rect(s, Inches(0.55), y, Inches(6.1), Inches(0.78), LIGHT,
         rounded=True, radius=0.18)
    rect(s, Inches(0.55), y, Inches(0.09), Inches(0.78), ORANGE)
    text(s, Inches(0.85), y + Inches(0.08), Inches(5.7), Inches(0.66),
         [(e, {"size": 14, "color": DARK})], line_spacing=0.95,
         anchor=MSO_ANCHOR.MIDDLE)
card(s, Inches(6.95), Inches(1.85), Inches(5.85), Inches(2.1),
     "Согласие сотрудника обязательно",
     "Переход на КЭДО возможен только с письменного согласия работника. "
     "Отказ — не основание для увольнения или отказа в приёме на работу.",
     accent=BLUE)
card(s, Inches(6.95), Inches(4.2), Inches(5.85), Inches(2.1),
     "Исключение для новичков",
     "Сотрудников, принятых после 31.12.2021 и не имеющих трудового стажа, "
     "можно переводить на КЭДО без отдельного согласия.", accent=GREEN)

# ============================================================ 10. Как перейти
s = prs.slides.add_slide(BLANK)
header(s, "Как перейти на КЭДО: 5 шагов",
       "Юридическая подготовка + настройка Битрикс24 — обычно 2–4 недели", 10)
steps10 = [
    ("Локальный нормативный акт", "Утвердите положение о КЭДО: система, "
     "порядок доступа, перечень документов и сроки."),
    ("Уведомление сотрудников", "Закон требует уведомить персонал о "
     "переходе на электронный документооборот до запуска."),
    ("Сбор согласий", "Получите письменные согласия — удобно собрать их "
     "прямо в Битрикс24."),
    ("Настройка Битрикс24 КЭДО", "Раздел «КЭДО + Госключ»: компания, "
     "представители, шаблоны, маршруты, интеграция с 1С."),
    ("Запуск и обучение", "Сотрудники устанавливают Госключ, подписывают "
     "первые документы — HR контролирует статусы."),
]
for i, (h, b) in enumerate(steps10):
    x = Inches(0.55 + (i % 3) * 4.18)
    y = Inches(1.8 + (i // 3) * 2.45)
    step_chip(s, x, y, Inches(3.95), Inches(2.25), i + 1, h, b,
              color=BLUE if i % 2 == 0 else CYAN)
rect(s, Inches(8.91), Inches(4.25), Inches(3.89), Inches(2.25), NAVY,
     rounded=True, radius=0.12)
text(s, Inches(9.2), Inches(4.5), Inches(3.4), Inches(1.9),
     [("Smart Business берёт на себя", {"size": 15, "bold": True,
                                        "color": CYAN}),
      ("шаблоны ЛНА, настройку портала, интеграцию с 1С и обучение "
       "команды.", {"size": 13, "color": WHITE})],
     line_spacing=1.1)

# ============================================================ 11. Выгоды и условия
s = prs.slides.add_slide(BLANK)
header(s, "Выгоды и условия подключения", None, 11)
num_card(s, Inches(0.55), Inches(1.75), Inches(2.95), Inches(2.3),
         "−80%", "времени на согласование и подписание кадровых документов",
         BLUE)
num_card(s, Inches(3.65), Inches(1.75), Inches(2.95), Inches(2.3),
         "30 ч", "в месяц экономии для каждого кадрового специалиста", CYAN)
num_card(s, Inches(6.75), Inches(1.75), Inches(2.95), Inches(2.3),
         "2 клика", "— и нужный документ найден в архиве даже через 5 лет",
         GREEN)
num_card(s, Inches(9.85), Inches(1.75), Inches(2.95), Inches(2.3),
         "100%", "юридическая значимость по 377-ФЗ и ТК РФ", ORANGE)
rect(s, Inches(0.55), Inches(4.35), Inches(12.25), Inches(2.15), LIGHT,
     rounded=True, radius=0.07)
text(s, Inches(0.95), Inches(4.6), Inches(11.5), Inches(0.5),
     [("Условия подключения", {"size": 17, "bold": True, "color": DARK})])
text(s, Inches(0.95), Inches(5.15), Inches(11.5), Inches(1.3),
     [("•  КЭДО входит в тарифы Битрикс24 «Профессиональный» и "
       "«Энтерпрайз» — без доплаты за пользователей и количество документов",
       {"size": 14, "color": GRAY}),
      ("•  Госключ для сотрудников бесплатен; отдельные лицензии на КЭДО "
       "покупать не нужно", {"size": 14, "color": GRAY}),
      ("•  Если вы уже на этих тарифах — инструмент достаточно настроить "
       "и запустить", {"size": 14, "color": GRAY})],
     line_spacing=1.2)

# ============================================================ 12. Финал / Smart Business
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, 0, SW, Inches(0.18), CYAN)
sb_logo(s, Inches(0.7), Inches(0.6), scale=1.2, dark_bg=True)
text(s, Inches(0.7), Inches(1.7), Inches(11.5), Inches(1.0),
     [("Внедрим Битрикс24 КЭДО под ключ", {"size": 40, "bold": True,
                                           "color": WHITE})])
text(s, Inches(0.7), Inches(2.75), Inches(11.5), Inches(0.6),
     [("Smart Business — официальный партнёр Битрикс24. Ростов-на-Дону, "
       "работаем с клиентами по всей России.",
       {"size": 16, "color": RGBColor(0xB9, 0xC9, 0xDE)})])
offers = [
    ("Аудит и план перехода", "Анализ кадровых процессов, перечень "
     "документов, шаблоны ЛНА и согласий."),
    ("Настройка и интеграции", "КЭДО + Госключ, шаблоны и маршруты, "
     "интеграция с 1С:ЗУП."),
    ("Обучение и поддержка", "Обучение HR и сотрудников, сопровождение "
     "после запуска."),
]
for i, (h, b) in enumerate(offers):
    x = Inches(0.7 + i * 4.1)
    rect(s, x, Inches(3.55), Inches(3.85), Inches(1.9),
         RGBColor(0x14, 0x2E, 0x52), rounded=True, radius=0.1)
    text(s, x + Inches(0.3), Inches(3.75), Inches(3.3), Inches(0.6),
         [(h, {"size": 16, "bold": True, "color": CYAN})])
    text(s, x + Inches(0.3), Inches(4.3), Inches(3.3), Inches(1.05),
         [(b, {"size": 13, "color": WHITE})], line_spacing=1.05)
rect(s, Inches(0.7), Inches(5.9), Inches(5.2), Inches(0.75), BLUE,
     rounded=True, radius=0.5)
text(s, Inches(0.7), Inches(5.9), Inches(5.2), Inches(0.75),
     [("Запросить демонстрацию КЭДО", {"size": 18, "bold": True,
                                       "color": WHITE})],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(6.2), Inches(5.95), Inches(6.5), Inches(0.7),
     [("bitrix24.ru/partners/partner/9471955", {"size": 14, "color": CYAN}),
      ], anchor=MSO_ANCHOR.MIDDLE)

OUT = "/home/user/-/presentation/Bitrix24_KEDO_Smart_Business.pptx"
prs.save(OUT)
print("saved:", OUT)
